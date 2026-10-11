import contextlib
import json
import logging
import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from api_falsa import ApiFalsa  # noqa: E402
from campos_permitidos import filtrar  # noqa: E402
from carregar_origem import (  # noqa: E402
    CHAVES, Contexto, PLANO_BULK, PLANO_SO_REST, RepositorioBanco, Tarefa, balancete_mensal, chave_origem,
    executar, meses_contabeis, opcoes_de_ambiente,
)
from importar_mapa_contas import importar, ler_planilha, texto_relatorio  # noqa: E402
from teste_carregar_origem import HOJE, OPCOES, TENANT, Origem, RepositorioFalso, cliente_para  # noqa: E402

for nome in ("cliente_origem", "carregar_origem"):
    logging.getLogger(nome).addHandler(logging.NullHandler())

TAREFA_BALANCETE = Tarefa("trial-balance", "balancete_mensal", "trial-balance")
NOME_TERCEIRO = "Fornecedor Pessoa Fisica"


def conta_balancete(empresa, numero, valor=100.0):
    return {
        "companyId": empresa, "reducedCode": numero, "accountGroup": "ASSETS",
        "account": f"1.1.1.01.{numero:04d}", "accountName": f"Conta {numero}", "accountType": "ANALYTICAL",
        "debit": valor, "credit": 0.0, "operatingBalance": {"value": valor, "type": "D"},
        "balances": {"previous": {"value": 0.0, "type": "D"}, "current": {"value": valor, "type": "D"}},
        "thirdParties": {"name": NOME_TERCEIRO, "code": 50, "debit": valor, "credit": 0.0},
        "campoNovo": "x",
    }


class OrigemContabil(Origem):
    """Origem com três obras em duas empresas, uma obra de outro tenant e balancete por empresa."""

    def __init__(self, contas_por_empresa=2):
        super().__init__()
        self.contas_por_empresa = contas_por_empresa
        self.valor = 100.0

    def rest(self):
        rest = super().rest()
        rest["cost-centers"] = [{"id": 101, "name": "Obra A", "idCompany": 11},
                                {"id": 102, "name": "Obra B", "idCompany": 12},
                                {"id": 103, "name": "Obra C", "idCompany": 11},
                                {"id": 999, "name": "Obra de outro tenant", "idCompany": 99}]
        rest["companies"] = [{"id": 11, "name": "SPE Aurora Ltda", "tradeName": "Aurora",
                              "cnpj": "11.111.111/0001-11", "address": "Rua A"}]
        rest["trial-balance"] = lambda consulta: [
            conta_balancete(int(consulta["companyId"]), n, self.valor) for n in range(1, self.contas_por_empresa + 1)]
        return rest


def balancetes(api):
    return [c["consulta"] for c in api.chamadas("rest", "trial-balance")]


def meses_pedidos(api, empresa):
    return sorted({c["initialPeriod"] for c in balancetes(api) if c["companyId"] == str(empresa)})


class TesteMesesContabeis(unittest.TestCase):
    def contexto(self, **ajustes):
        return Contexto(cliente=None, hoje=HOJE, obras=[], **ajustes)

    def test_primeira_carga_vem_do_inicio_contabil_ate_o_mes_corrente(self):
        meses = meses_contabeis(self.contexto(inicio_contabil=date(2026, 5, 1)), None, 11)
        self.assertEqual(meses, [date(2026, m, 1) for m in range(5, 10)])

    def test_sem_inicio_contabil_usa_a_janela_inicial(self):
        meses = meses_contabeis(self.contexto(janela_dias=400), None, 11)
        self.assertEqual((meses[0], meses[-1], len(meses)), (date(2025, 8, 1), date(2026, 9, 1), 14))

    def test_incremental_reve_os_tres_ultimos_meses(self):
        ctx = self.contexto(inicio_contabil=date(2024, 1, 1), empresas_com_balancete={11})
        self.assertEqual(meses_contabeis(ctx, HOJE - timedelta(days=1), 11),
                         [date(2026, 7, 1), date(2026, 8, 1), date(2026, 9, 1)])

    def test_marca_antiga_reve_desde_o_mes_da_marca(self):
        ctx = self.contexto(inicio_contabil=date(2024, 1, 1), empresas_com_balancete={11})
        meses = meses_contabeis(ctx, date(2026, 3, 15), 11)
        self.assertEqual((meses[0], meses[-1]), (date(2026, 1, 1), date(2026, 9, 1)))

    def test_revisao_nao_volta_antes_do_inicio_contabil(self):
        ctx = self.contexto(inicio_contabil=date(2026, 8, 1), empresas_com_balancete={11})
        self.assertEqual(meses_contabeis(ctx, HOJE, 11), [date(2026, 8, 1), date(2026, 9, 1)])

    def test_virada_de_ano(self):
        ctx = Contexto(cliente=None, hoje=date(2027, 1, 10), obras=[], empresas_com_balancete={11})
        self.assertEqual(meses_contabeis(ctx, date(2027, 1, 9), 11),
                         [date(2026, 11, 1), date(2026, 12, 1), date(2027, 1, 1)])


class TesteCargaContabil(unittest.TestCase):
    def rodar(self, origem, repo, hoje=HOJE, **opcoes):
        with ApiFalsa(rest=origem.rest(), bulk=origem.bulk()) as api:
            relatorio = executar(cliente_para(api), repo, TENANT, hoje, **{**OPCOES, **opcoes})
        return relatorio, api

    def test_tarefas_nos_dois_caminhos_com_lista_e_chave(self):
        for plano in (PLANO_BULK, PLANO_SO_REST):
            self.assertIn("companies", {t.endpoint for t in plano})
            self.assertIn("trial-balance", {t.endpoint for t in plano})
        self.assertEqual(CHAVES["trial-balance"].campos,
                         ("companyId", "account", "initialPeriod", "desconsiderClosingBalance"))
        self.assertEqual(CHAVES["companies"].campos, ("id",))

    def test_parametros_exatos_das_duas_passagens(self):
        origem, repo = OrigemContabil(), RepositorioFalso(obras=(101,))
        relatorio, api = self.rodar(origem, repo, inicio_contabil=date(2026, 2, 1))
        self.assertEqual(relatorio.falhas, [])
        fevereiro = [c for c in balancetes(api) if c["initialPeriod"] == "2026-02-01"]
        esperado = {"companyId": "11", "initialPeriod": "2026-02-01", "finalPeriod": "2026-02-28",
                    "reportType": "CORPORATE", "onlyAccountsWithMovements": "false", "thirdParties": "false",
                    "limit": "200", "offset": "0"}
        self.assertEqual(sorted(fevereiro, key=lambda c: c["desconsiderClosingBalance"]),
                         [{**esperado, "desconsiderClosingBalance": "false"},
                          {**esperado, "desconsiderClosingBalance": "true"}])
        self.assertEqual(len(balancetes(api)), 8 * 2)

    def test_empresas_das_obras_do_tenant_mais_a_variavel(self):
        origem, repo = OrigemContabil(), RepositorioFalso(obras=(101, 103))
        _, api = self.rodar(origem, repo, inicio_contabil=date(2026, 9, 1), empresas_contabeis=(7,))
        self.assertEqual({c["companyId"] for c in balancetes(api)}, {"11", "7"})
        self.assertEqual(len(api.chamadas("rest", "cost-centers")), 2, "cadastro e uma leitura para as empresas")

    def test_registro_gravado_com_mes_passagem_e_sem_terceiros(self):
        origem, repo = OrigemContabil(), RepositorioFalso(obras=(101,))
        self.rodar(origem, repo, inicio_contabil=date(2026, 9, 1))
        registros = repo.payloads("trial-balance")
        self.assertEqual(len(registros), 2 * 2)
        primeiro = sorted(registros, key=lambda r: (r["account"], r["desconsiderClosingBalance"]))[0]
        self.assertEqual(primeiro, {
            "companyId": 11, "reducedCode": 1, "accountGroup": "ASSETS", "account": "1.1.1.01.0001",
            "accountName": "Conta 1", "accountType": "ANALYTICAL", "debit": 100.0, "credit": 0.0,
            "operatingBalance": {"value": 100.0, "type": "D"},
            "balances": {"previous": {"value": 0.0, "type": "D"}, "current": {"value": 100.0, "type": "D"}},
            "initialPeriod": "2026-09-01", "desconsiderClosingBalance": False,
        })
        self.assertNotIn(NOME_TERCEIRO, json.dumps(registros))
        self.assertEqual(chave_origem("trial-balance", primeiro), "11|1.1.1.01.0001|2026-09-01|False")
        empresas = repo.payloads("companies")
        self.assertEqual(empresas, [{"id": 11, "name": "SPE Aurora Ltda", "tradeName": "Aurora",
                                     "cnpj": "11.111.111/0001-11"}])

    def test_terceiros_nunca_passam_no_filtro(self):
        self.assertNotIn("thirdParties", filtrar("trial-balance", conta_balancete(11, 1)))

    def test_incremental_pede_tres_meses_e_empresa_nova_vem_inteira(self):
        origem, repo = OrigemContabil(), RepositorioFalso(obras=(101,))
        _, api = self.rodar(origem, repo, inicio_contabil=date(2026, 1, 1))
        self.assertEqual(len(meses_pedidos(api, 11)), 9)
        _, api = self.rodar(origem, repo, hoje=HOJE + timedelta(days=1), inicio_contabil=date(2026, 1, 1),
                            empresas_contabeis=(7,))
        self.assertEqual(meses_pedidos(api, 11), ["2026-07-01", "2026-08-01", "2026-09-01"])
        self.assertEqual(len(meses_pedidos(api, 7)), 9)

    def test_reler_mes_substitui_sem_duplicar(self):
        origem, repo = OrigemContabil(), RepositorioFalso(obras=(101,))
        self.rodar(origem, repo, inicio_contabil=date(2026, 8, 1))
        self.assertEqual(repo.contagem()["trial-balance"], 2 * 2 * 2)
        origem.valor = 250.0
        self.rodar(origem, repo, hoje=HOJE + timedelta(days=1), inicio_contabil=date(2026, 8, 1))
        registros = repo.payloads("trial-balance")
        self.assertEqual(len(registros), 2 * 2 * 2)
        self.assertEqual({r["debit"] for r in registros}, {250.0})

    def test_varias_paginas_sem_repeticao_passam(self):
        origem, repo = OrigemContabil(contas_por_empresa=450), RepositorioFalso(obras=(101,))
        relatorio, api = self.rodar(origem, repo, inicio_contabil=date(2026, 9, 1))
        self.assertEqual(relatorio.falhas, [])
        self.assertEqual(len(balancetes(api)), 2 * 3)
        self.assertEqual(repo.contagem()["trial-balance"], 450 * 2)

    def test_pagina_repetida_falha_sem_gravar(self):
        origem, repo = OrigemContabil(), RepositorioFalso(obras=(101,))
        contas = [conta_balancete(11, n) for n in range(1, 201)]

        def ignora_offset(metodo, consulta, corpo, autorizacao):
            return 200, {"resultSetMetadata": {"count": 600, "offset": 0, "limit": 200}, "results": contas}

        with ApiFalsa(rest=origem.rest(), bulk=origem.bulk(), rotas={"trial-balance": ignora_offset}) as api:
            relatorio = executar(cliente_para(api), repo, TENANT, HOJE, **OPCOES, inicio_contabil=date(2026, 9, 1))
        falhas = dict(relatorio.falhas)
        self.assertIn("página repetida na empresa 11, mês 2026-09, com zeramento", falhas["trial-balance"])
        self.assertEqual(len(api.chamadas("rest", "trial-balance")), 2)
        self.assertEqual(repo.contagem()["trial-balance"], 0)
        self.assertNotIn((TENANT, "trial-balance"), repo.marcas)
        self.assertIn((TENANT, "companies"), repo.marcas)

    def test_estrategia_sem_obra_e_sem_variavel_nao_chama(self):
        with ApiFalsa(rest=OrigemContabil().rest()) as api:
            ctx = Contexto(cliente=cliente_para(api), hoje=HOJE, obras=[])
            self.assertEqual(list(balancete_mensal(ctx, TAREFA_BALANCETE, None)), [])
            self.assertEqual(api.requisicoes, [])


class TesteAmbienteEStaging(unittest.TestCase):
    def test_variaveis_novas(self):
        opcoes = opcoes_de_ambiente({"ORIGEM_INICIO_CONTABIL": "2023-01", "ORIGEM_EMPRESAS_CONTABEIS": " 12, 1,12"})
        self.assertEqual((opcoes["inicio_contabil"], opcoes["empresas_contabeis"]), (date(2023, 1, 1), (1, 12)))
        vazio = opcoes_de_ambiente({})
        self.assertEqual((vazio["inicio_contabil"], vazio["empresas_contabeis"]), (None, ()))

    def test_variaveis_invalidas_dizem_o_formato(self):
        with self.assertRaisesRegex(ValueError, "aaaa-mm"):
            opcoes_de_ambiente({"ORIGEM_INICIO_CONTABIL": "01/2023"})
        with self.assertRaisesRegex(ValueError, "separados por vírgula"):
            opcoes_de_ambiente({"ORIGEM_EMPRESAS_CONTABEIS": "1;2"})

    def test_recarga_chama_o_staging_contabil_depois_dos_outros(self):
        comandos = []

        class Cursor:
            def execute(self, sql, parametros=None):
                comandos.append(sql)

        class Conexao:
            def transaction(self):
                return contextlib.nullcontext()

            def cursor(self):
                return contextlib.nullcontext(Cursor())

        RepositorioBanco(Conexao()).recarregar(TENANT)
        self.assertEqual(comandos[-1], "select staging.recarregar_contabil(%s)")
        self.assertEqual(len(comandos), 4)


CATALOGO = [
    ("bancos", "Bancos c/movimento", "1"),
    ("fornecedores", "Fornecedores e Credores Diversos", "2"),
    ("servicos_terceiros_despesa", "Serviços de Terceiros", "4"),
    ("servicos_terceiros_obra", "Serviços de Terceiros", "5"),
    ("custo_unidades_vendidas", "Custo das Unidades Vendidas", "4"),
]
NOME_PESSOA = "Mutuo Socio Exemplo"


class BancoMapaFalso:
    def __init__(self, vigente=None):
        self.vigente = dict(vigente or {})
        self.gravados = []
        self.transacoes = 0

    @contextlib.contextmanager
    def transacao(self):
        self.transacoes += 1
        yield

    def ler_catalogo(self):
        return CATALOGO

    def ler_vigente(self, tenant):
        return dict(self.vigente)

    def gravar(self, tenant, itens):
        self.gravados.extend(itens)


class TesteImportarMapaContas(unittest.TestCase):
    def setUp(self):
        from openpyxl import Workbook

        self.pasta = tempfile.TemporaryDirectory()
        self.arquivo = Path(self.pasta.name) / "plano.xlsx"
        livro = Workbook()
        livro.active.title = "Outra"
        aba = livro.create_sheet("Plano de Contas")
        aba.append(["Reduzido", "Conta", "Nome", "Tipo", "Natureza", "Linha Dasboard"])
        linhas = [
            (1, "1.1.1.02.0001", "Banco A", "  bancos C/MOVIMENTO "),
            (2, "2.1.1.01.0001", "Fornecedores", "Fornecedores e Credores Diversos"),
            (3, "4.2.1.01.0003", "Servicos", "Serviços de Terceiros"),
            (4, "5.1.1.01.0003", "Servicos", "serviços  de terceiros"),
            (5, "4.1.1.01.0001", "CMV", "Custo das Unidades Vendidas"),
            (6, "2.2.1.01.0009", NOME_PESSOA, "Mútuo de Sócio Inexistente"),
            (7, "2.2.1.01.0010", NOME_PESSOA, "Mútuo de Sócio Inexistente"),
            (8, "3.1.1.01.0001", "Sem linha", None),
        ]
        for reduzido, conta, nome, rotulo in linhas:
            aba.append([reduzido, conta, nome, "A", "D", rotulo])
        aba.append([])
        livro.save(self.arquivo)

    def tearDown(self):
        self.pasta.cleanup()

    def test_leitura_ignora_cabecalho_linha_vazia_e_sem_rotulo(self):
        pares = ler_planilha(self.arquivo)
        self.assertEqual(len(pares), 7)
        self.assertEqual(pares[0], ("1.1.1.02.0001", "bancos C/MOVIMENTO"))

    def test_casa_rotulo_e_grupo_e_grava_so_o_que_mudou(self):
        banco = BancoMapaFalso(vigente={"2.1.1.01.0001": "fornecedores", "4.1.1.01.0001": "outros_custos"})
        plano, ambiguos = importar(banco, TENANT, ler_planilha(self.arquivo), simular=False)
        self.assertEqual(banco.transacoes, 1)
        self.assertEqual(ambiguos, set())
        self.assertEqual(banco.gravados, [
            ("1.1.1.02.0001", "bancos"),
            ("4.1.1.01.0001", "custo_unidades_vendidas"),
            ("4.2.1.01.0003", "servicos_terceiros_despesa"),
            ("5.1.1.01.0003", "servicos_terceiros_obra"),
        ])
        self.assertEqual(plano.iguais, 1)
        self.assertEqual(plano.sem_correspondencia,
                         {("Mútuo de Sócio Inexistente", "2"): ["2.2.1.01.0009", "2.2.1.01.0010"]})

    def test_simular_nao_grava_e_relatorio_sem_nome_de_conta(self):
        banco = BancoMapaFalso()
        plano, ambiguos = importar(banco, TENANT, ler_planilha(self.arquivo), simular=True)
        self.assertEqual(banco.gravados, [])
        texto = texto_relatorio(plano, ambiguos, simular=True)
        self.assertIn("contas mapeadas (seriam gravadas): 5", texto)
        self.assertIn("contas já iguais ao mapa vigente: 0", texto)
        self.assertIn("Mútuo de Sócio Inexistente (grupo 2): 2 contas: 2.2.1.01.0009, 2.2.1.01.0010", texto)
        for nome in (NOME_PESSOA, "Banco A", "CMV"):
            self.assertNotIn(nome, texto)

    def test_conta_repetida_com_linhas_diferentes_nao_grava(self):
        pares = [("4.2.1.01.0003", "Serviços de Terceiros"), ("4.2.1.01.0003", "Custo das Unidades Vendidas"),
                 ("1.1.1.02.0001", "Bancos c/movimento")]
        banco = BancoMapaFalso()
        plano, _ = importar(banco, TENANT, pares, simular=False)
        self.assertEqual(plano.conflitos, ["4.2.1.01.0003"])
        self.assertEqual(banco.gravados, [("1.1.1.02.0001", "bancos")])

    def test_aba_ausente_diz_o_que_falta(self):
        from openpyxl import Workbook

        outro = Path(self.pasta.name) / "sem_aba.xlsx"
        Workbook().save(outro)
        with self.assertRaisesRegex(ValueError, "Plano de Contas"):
            ler_planilha(outro)


if __name__ == "__main__":
    unittest.main()
